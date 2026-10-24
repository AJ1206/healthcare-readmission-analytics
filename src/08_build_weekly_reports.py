"""Step 8 - Build the six weekly internship submissions (Word) from the project's real outputs.

Weeks 1-5 are self-contained documents written to each week's brief; Week 6 is the integrated report
produced by 07_build_report.py. Run after steps 1-7.
"""
from __future__ import annotations

import shutil
import sys
from datetime import date
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import weekly_figures as wf  # noqa: E402
from common import FIG_DIR, ROOT  # noqa: E402
from docx_helpers import ACCENT, MUTED, Report  # noqa: E402
from docx.shared import Pt  # noqa: E402
from report_data import *  # noqa: E402,F401,F403

OUT_DIR = ROOT / "report" / "weekly"
OUT_DIR.mkdir(parents=True, exist_ok=True)
FIGS = wf.all_figures(RISK_POINTS)
F = lambda name: FIG_DIR / name  # noqa: E731


# ------------------------------------------------------------------ shared building blocks
def new_doc(week: int, title: str, subtitle: str) -> Report:
    r = Report()
    r.page_numbers(f"Week {week} · {title}")
    for _ in range(5):
        r.p("")
    r.p(f"SIX-WEEK HEALTHCARE DATA ANALYTICS INTERNSHIP · WEEK {week} SUBMISSION", size=10, color=MUTED, align="center")
    t = r.p(title, size=24, align="center")
    t.runs[0].bold = True
    t.runs[0].font.color.rgb = ACCENT
    r.p(subtitle, size=12.5, align="center")
    r.p("")
    r.p("Project: 30-day readmission and resource use among diabetic inpatients in 130 US hospitals "
        "(UCI Diabetes 130-US Hospitals dataset, 1999–2008)", size=11, italic=True, align="center")
    for _ in range(6):
        r.p("")
    for lead, val in [("Prepared by: ", AUTHOR), ("Internship organisation: ", "[Organisation name]"),
                      ("Supervisor: ", "[Supervisor name]"), ("Date: ", date.today().strftime("%d %B %Y")),
                      ("Project repository: ", REPO_URL)]:
        r.p(val, bold_lead=lead, align="center", space_after=2)
    r.page_break()
    r.toc("Contents", 'TOC \\o "1-2" \\h \\z \\u', "Right-click and choose Update Field (or press F9) to build the contents.")
    r.page_break()
    return r


def references(r: Report, keys: list[str]):
    r.h("References", page_break=True)
    for ref in sorted(REFS[k] for k in keys):
        para = r.p(ref, size=9.5, space_after=6)
        para.paragraph_format.left_indent = Pt(24)
        para.paragraph_format.first_line_indent = Pt(-24)


def how_to_read(r: Report, text: str):
    r.p(text, italic=True, color=MUTED, size=9.5, align="justify")


def risk_table(r: Report, caption="Risk register"):
    r.table(RISKS.astype(str), caption, widths_cm=[1.0, 3.6, 1.4, 1.2, 1.0, 3.2, 3.0, 2.2], font_size=7.5)


def kpi_table(r: Report, caption="KPI catalogue"):
    r.table(KPIS, caption, widths_cm=[4.2, 6.6, 2.4, 3.4], font_size=8)


def save(r: Report, name: str) -> Path:
    path = OUT_DIR / name
    r.save(path)
    print(f"Saved {path.name}: {r.fig_n} figures, {r.tab_n} tables")
    return path


# ================================================================== WEEK 1
def week1() -> Path:
    r = new_doc(1, "Healthcare Data Analysis Strategic Plan", "Planning, organisation and roadmap for a diabetic readmission analytics project")
    how_to_read(r, "This plan is self-contained: it defines the problem, the analytical strategy, the metrics, the "
                   "roadmap and the risk controls needed to run the project. Baseline values quoted in the KPI and "
                   "scope sections come from the project's own data, because the plan was subsequently executed "
                   f"(code and outputs: {REPO_URL}).")

    r.h("1. Executive Summary")
    r.p("Unplanned readmissions within 30 days of discharge are a recognised marker of care quality and a major "
        "cost for hospitals; in the US Medicare population about one in five discharged patients was rehospitalised "
        "within 30 days (Jencks et al., 2009). Patients with diabetes are a large, high-risk group. This plan sets "
        "out how a hospital analytics team would use routinely collected inpatient data to understand who is "
        "readmitted, how inpatient resources are used and whether readmission risk can be predicted early enough to "
        "act.", align="justify")
    r.p("Vision: ", bold_lead="")
    r.doc.paragraphs[-1].runs[0].bold = True
    r.p("To give clinical and operational leaders a transparent, reproducible and ethically governed evidence base "
        "on diabetic inpatient readmissions that can be refreshed with local data and used to target "
        "transition-of-care resources.", italic=True, align="justify")
    r.bullets([("Objective 1 – Baseline. ", "Measure 30-day readmission, length of stay and care-process KPIs with "
                "95% confidence intervals."),
               ("Objective 2 – Drivers. ", "Identify patient, clinical and discharge factors associated with early "
                "readmission, using pre-specified hypotheses and effect sizes."),
               ("Objective 3 – Prediction. ", "Build and honestly validate a risk model that beats a simple rule."),
               ("Objective 4 – Communication. ", "Deliver a dashboard and a report that non-technical stakeholders "
                "can act on."),
               ("Objective 5 – Governance. ", "Meet privacy, ethics and reproducibility standards throughout.")])
    r.p("The project runs for six weeks (30–35 hours per week, 195 hours in total) with six milestone reviews. "
        "Its main deliverables are a quality-assured dataset, a statistical analysis, a validated predictive model, "
        "an interactive dashboard and an integrated report with evidence-based recommendations.", align="justify")

    r.h("2. The Role of Strategic Planning in Healthcare Data Analysis", page_break=True)
    r.p("Healthcare data analysis differs from analytics in most other sectors in four ways that make up-front "
        "planning essential:", align="justify")
    r.bullets([("Complex, imperfect data. ", "Clinical and administrative data are collected for care and billing, "
                "not research. They contain coding conventions, missing values and repeated measurements that must be "
                "understood before analysis (Kahn et al., 2016; Weiskopf & Weng, 2013)."),
               ("High stakes. ", "Findings can influence patient care. Errors such as data leakage or confusing "
                "association with causation can cause harm, so methods must be fixed before results are seen."),
               ("Regulation and ethics. ", "Personal health data are protected by law (for example India's Digital "
                "Personal Data Protection Act 2023 and Rules 2025; HIPAA in the US). Access, consent and retention "
                "must be planned, not retrofitted (MeitY, 2025; HHS, 2012)."),
               ("Many stakeholders. ", "Clinicians, managers, data protection officers and patients have different "
                "questions; a plan aligns them on what will be delivered and how success is judged.")])
    r.p("Big-data analytics can improve care quality and reduce costs only when it is organised around clear "
        "questions, governance and the right data (Raghupathi & Raghupathi, 2014). This plan follows the "
        "Cross-Industry Standard Process for Data Mining (CRISP-DM), whose six phases — business understanding, "
        "data understanding, data preparation, modelling, evaluation and deployment — provide a widely used "
        "structure for analytics projects (Wirth & Hipp, 2000). It is adapted here with explicit quality-assurance, "
        "ethics and communication stages (Figure 1).", align="justify")
    r.figure(FIGS["workflow"], "Analytical workflow used in this plan (adapted from CRISP-DM)",
             "Each stage has defined inputs, outputs and acceptance criteria (Section 7), and feedback from monitoring "
             "returns to planning.")

    r.h("3. Problem Statement and Business Case", page_break=True)
    r.p("A hospital quality team needs to know how often diabetic inpatients return within 30 days, which factors "
        "are associated with early readmission, where inpatient bed-days are concentrated, and whether discharge data "
        "can rank patients by risk well enough to prioritise follow-up. Without this evidence, transition-of-care "
        "resources are spread evenly instead of targeted where risk is highest.", align="justify")
    r.table(pd.DataFrame([
        ["Fewer avoidable readmissions", "Structured discharge programmes reduced readmissions in randomised trials (Jack et al., 2009; Naylor et al., 1999; Leppin et al., 2014)"],
        ["Freed bed capacity", "Each avoided readmission releases an inpatient bed for a new admission"],
        ["Financial and regulatory exposure", "Readmission performance is publicly reported and payment-linked in several systems (CMS, 2026)"],
        ["Better diabetes care", "Inconsistent inpatient diabetes management is linked to readmission (Strack et al., 2014)"],
        ["Reusable capability", "A scripted pipeline turns a one-off study into a monitoring tool"],
    ], columns=["Benefit", "Rationale and evidence"]), "Business case", widths_cm=[4.4, 12.2])

    r.h("4. Project Scope", page_break=True)
    r.table(pd.DataFrame([
        ["In scope", "Patient-level analysis of diabetic inpatient encounters; 30-day readmission; length of stay, bed-days, medications and laboratory use; care-process indicators (HbA1c testing); descriptive, inferential and predictive analysis; dashboard; governance and ethics review"],
        ["Out of scope", "Causal effect estimation; cost analysis; real-time deployment into clinical systems; individual patient decisions"],
        ["Assumptions", "A suitable public dataset with a patient-level readmission outcome exists and is licensed for reuse; readmission labels are coded correctly; standard laptop computing is sufficient"],
        ["Constraints", "Six weeks; public data only (no credentialed clinical databases); open-source tools; a single analyst with supervisor review"],
        ["Dependencies", "Dataset availability; supervisor reviews at each milestone; clinician input on variable definitions"],
    ], columns=["Element", "Definition"]), "Scope statement", widths_cm=[3.0, 13.6])
    r.h("4.1 Objectives (SMART)", 2)
    r.table(pd.DataFrame([
        ["O1", "Acquire and document a licensed healthcare dataset", "Source, licence, checksum and schema recorded", "End of week 2"],
        ["O2", "Deliver an analysis-ready dataset", "100% of validation rules pass; issue log complete", "End of week 2"],
        ["O3", "Quantify readmission and resource-use KPIs", "≥ 10 KPIs with formulas and 95% CIs", "End of week 3"],
        ["O4", "Test pre-specified hypotheses", "Effect sizes reported; multiplicity controlled", "Mid week 4"],
        ["O5", "Build and validate a risk model", "Beats a simple baseline on an untouched test set; calibration and subgroups reported", "End of week 4"],
        ["O6", "Provide an interactive dashboard", "Filters, KPI cards, confidence intervals, small-sample suppression", "End of week 3"],
        ["O7", "Produce traceable recommendations", "Each linked to a result or cited source", "End of week 5"],
    ], columns=["ID", "Objective", "Measure of success", "Deadline"]), "Objectives", widths_cm=[1.0, 5.6, 7.0, 3.0])
    r.h("4.2 Key questions and hypotheses", 2)
    r.table(pd.DataFrame([
        ["Q1", "What is the 30-day readmission rate, and how does it vary across patient groups?", "Descriptive"],
        ["Q2", "Is prior inpatient use associated with readmission?", "H1: ≥ 1 prior-year admission → higher rate"],
        ["Q3", "Does discharge destination matter?", "H2: discharge to a care facility → higher rate"],
        ["Q4", "Is inpatient HbA1c measurement associated with readmission?", "H3: measurement → lower rate after adjustment; varies by diagnosis"],
        ["Q5", "Are length of stay and age associated with readmission?", "H4: longer stay and older age → higher rate"],
        ["Q6", "Does resource use differ by primary diagnosis?", "H5: length of stay differs across diagnosis groups"],
        ["Q7", "Can discharge data rank risk better than a simple rule?", "H6: multivariable model beats a prior-admission rule (ROC-AUC)"],
    ], columns=["ID", "Key question", "Hypothesis / type"]), "Key questions and testable hypotheses", widths_cm=[1.0, 8.6, 7.0])

    r.h("5. Stakeholders and Governance", page_break=True)
    r.table(pd.DataFrame([
        ["Executive leadership", "High", "High", "Monthly one-page KPI brief"],
        ["CMO / quality committee", "High", "High", "Milestone reviews; sign-off of findings"],
        ["Diabetes / endocrinology team", "Medium", "High", "Clinical review of HbA1c analysis"],
        ["Hospitalists, nurses, case managers", "Medium", "High", "Co-design of dashboard and pilot"],
        ["Health information management", "Medium", "Medium", "Data-quality issue log"],
        ["IT / data engineering", "Medium", "Medium", "Pipeline hand-over"],
        ["Data protection officer", "High", "Medium", "Approval before any identifiable data is used"],
        ["Patients and caregivers", "Low", "High", "Plain-language summary via patient advisory group"],
    ], columns=["Stakeholder", "Influence", "Interest", "Engagement approach"]), "Stakeholder map", widths_cm=[5.2, 2.0, 2.0, 7.4])
    r.h("5.1 Responsibilities (RACI)", 2)
    r.table(pd.DataFrame([
        ["Project plan and charter", "R", "A", "C", "C", "I"],
        ["Data acquisition and QA", "R", "A", "I", "C", "C"],
        ["Statistical analysis", "R", "A", "C", "I", "I"],
        ["Predictive model", "R", "A", "C", "C", "C"],
        ["Dashboard and visualisation", "R", "A", "C", "I", "I"],
        ["Recommendations", "R", "C", "A", "I", "I"],
        ["Privacy and ethics review", "C", "R", "I", "I", "A"],
    ], columns=["Activity", "Analyst (intern)", "Supervisor", "Clinical lead", "IT", "Data protection officer"]),
        "RACI matrix (R = responsible, A = accountable, C = consulted, I = informed)",
        widths_cm=[5.0, 2.3, 2.3, 2.3, 1.6, 3.1], font_size=8.5)
    r.h("5.2 Governance procedures", 2)
    r.bullets(["Weekly milestone review against acceptance criteria; decisions recorded in a change log.",
               "Analytical peer review: a second analyst re-runs the pipeline and checks key numbers.",
               "Clinical review of variable definitions (diagnosis grouping, discharge grouping).",
               "Privacy checkpoint before any non-public data is used (data-protection impact assessment).",
               "Model governance: no clinical use without validation, an accountable owner and monitoring."])

    r.h("6. Data Strategy", page_break=True)
    r.p("The data strategy prioritises sources that are accessible, licensed for reuse, well documented and contain a "
        "patient-level readmission outcome. Candidate sources are scored against nine criteria (detailed in the "
        "Week 2 plan). The selected source is the UCI Diabetes 130-US Hospitals dataset: 101,766 inpatient encounters "
        "of patients with diabetes (1999–2008), licensed CC BY 4.0, with a documented readmission outcome "
        "(Clore et al., 2014; Strack et al., 2014). Alternatives held in reserve are the CMS Hospital Readmissions "
        "Reduction Program data (hospital-level) and, for longer projects, credentialed databases such as MIMIC-IV.",
        align="justify")

    r.h("7. Step-by-Step Analytical Strategy", page_break=True)
    steps = [
        ("7.1 Data acquisition", ["Confirm access and licence; record source URL, DOI and access date.",
                                  "Download programmatically; store the untouched original and its SHA-256 checksum.",
                                  "Extract only expected files; never edit raw data."],
         "Raw data + provenance file", "Checksum recorded; licence permits use"),
        ("7.2 Data profiling and cleaning", ["Profile types, missingness, duplicates, invalid codes and outliers.",
                                            "Apply validation rules (ranges, domains, cross-field logic, uniqueness).",
                                            "Clean with documented, justified rules; log every transformation.",
                                            "Define the analysis cohort (one row per patient; exclude deaths and hospice)."],
         "Clean dataset, data dictionary, QA report", "All rules pass; issue log complete"),
        ("7.3 Exploratory analysis", ["Describe distributions and group differences.",
                                     "Compute KPIs with 95% Wilson confidence intervals.",
                                     "Produce publication-quality figures with interpretations."],
         "Figure set and descriptive tables", "Every chart titled, labelled and interpreted"),
        ("7.4 Statistical analysis", ["Chi-square, Mann–Whitney and Kruskal–Wallis tests with effect sizes.",
                                     "Holm correction for multiple comparisons.",
                                     "Multivariable logistic regression for adjusted associations."],
         "Test results; adjusted odds ratios", "Assumptions checked; effect sizes reported"),
        ("7.5 Predictive modelling", ["Patient-level 60/20/20 train/validation/test split.",
                                     "Baselines, then logistic regression, random forest and gradient boosting.",
                                     "Tune with cross-validation on training data only; evaluate once on test data.",
                                     "Report discrimination, calibration and subgroup performance."],
         "Validated model + model card", "Beats baseline; no leakage; test set used once"),
        ("7.6 Visualisation and reporting", ["Interactive dashboard with filters and suppression of small groups.",
                                            "Analytical report with traceable recommendations."],
         "Dashboard; report", "Stakeholder review passed"),
    ]
    for title, acts, output, accept in steps:
        r.h(title, 2)
        r.bullets(acts)
        r.p(f"{output}. ", bold_lead="Output: ")
        r.p(f"{accept}.", bold_lead="Acceptance criterion: ")
    r.h("7.7 Tools and resources", 2)
    r.table(pd.DataFrame([
        ["Python (pandas, NumPy)", "Data loading, profiling, cleaning (McKinney, 2010)"],
        ["SciPy, statsmodels", "Hypothesis tests, confidence intervals, regression"],
        ["scikit-learn", "Pipelines, model training, cross-validation, metrics (Pedregosa et al., 2011)"],
        ["Matplotlib, seaborn", "Static publication-quality charts"],
        ["Power BI / HTML dashboard", "Interactive dashboard for stakeholders"],
        ["SQL", "Validation rules and extracts when working with a hospital database"],
        ["Git and GitHub", "Version control and reproducibility"],
        ["Microsoft Word", "Reports and stakeholder documents"],
    ], columns=["Tool", "Purpose"]), "Tools and their purpose", widths_cm=[5.0, 11.6])

    r.h("8. Key Metrics (KPI Catalogue)", page_break=True)
    r.p("Each KPI has a fixed formula so that it can be recalculated identically every reporting period. Baseline "
        "values are those measured when the plan was executed on the selected dataset (analysis cohort of "
        f"{num(ov['n'])} patients).", align="justify")
    kpi_table(r)

    r.h("9. Deliverables", page_break=True)
    r.table(pd.DataFrame([
        ["D1", "Strategic plan (this document)", "Week 1", "Supervisor"],
        ["D2", "Data collection and QA plan; clean dataset; data dictionary; QA report", "Week 2", "Supervisor, HIM"],
        ["D3", "Visualisation plan; figure set; interactive dashboard", "Week 3", "Quality committee"],
        ["D4", "Statistical and predictive modelling framework; model and evaluation", "Week 4", "Clinical lead"],
        ["D5", "Analytical report with recommendations", "Week 5", "CMO, executives"],
        ["D6", "Integrated final report, code repository, self-assessment", "Week 6", "All stakeholders"],
    ], columns=["ID", "Deliverable", "Due", "Audience"]), "Deliverables register", widths_cm=[1.0, 9.6, 2.2, 3.8])

    r.h("10. Roadmap and Timeline", page_break=True)
    r.figure(FIGS["gantt"], "Six-week roadmap", "Phases overlap slightly so that, for example, dashboard design starts "
             "while exploratory figures are being finalised. Each week ends with a milestone review.")
    road = pd.DataFrame([
        ["1", "Strategic planning", "Problem definition, stakeholder interviews, objectives, KPIs, risk register, roadmap", "32", "Supervisor availability", "M1 Charter approved"],
        ["2", "Data acquisition & QA", "Source comparison, download, profiling, validation, cleaning, data dictionary", "34", "Dataset access", "M2 Clean data + QA report"],
        ["3", "EDA & visualisation", "Descriptive statistics, figures, dashboard build, stakeholder feedback", "34", "Clean data", "M3 Figures + dashboard"],
        ["4", "Statistics & modelling", "Hypothesis tests, regression, model training and validation", "35", "EDA findings", "M4 Model validated"],
        ["5", "Interpretation & reporting", "Synthesis, recommendations, analytical report", "30", "Weeks 2–4 results", "M5 Draft report"],
        ["6", "Integration & reflection", "Integrated report, packaging, self-assessment, final review", "30", "All outputs", "M6 Final submission"],
    ], columns=["Week", "Phase", "Main activities", "Hours", "Dependency", "Milestone"])
    r.table(road, "Weekly plan with hours, dependencies and milestones", widths_cm=[1.0, 2.8, 6.4, 1.2, 2.6, 2.6], font_size=8)
    r.h("10.1 Milestone definitions", 2)
    r.table(pd.DataFrame([
        ["M1", "Charter approved", "Objectives, scope, KPIs and risks signed off by supervisor"],
        ["M2", "Clean data + QA report", "All validation rules pass; issue log and data dictionary complete"],
        ["M3", "Figures + dashboard", "Every figure interpreted; dashboard reviewed by one non-technical user"],
        ["M4", "Model validated", "Test-set metrics, calibration and subgroup audit reported; beats baseline"],
        ["M5", "Draft report", "All sections drafted; every recommendation traceable to evidence"],
        ["M6", "Final submission", "Integrated report, repository and self-assessment complete; numbers consistent"],
    ], columns=["ID", "Milestone", "Completion criteria"]), "Milestone definitions", widths_cm=[1.0, 4.0, 11.6])

    r.h("11. Resource Allocation", page_break=True)
    r.table(pd.DataFrame([
        ["Planning and stakeholder work", "32", "16%"], ["Data acquisition, QA and cleaning", "34", "17%"],
        ["Exploration and visualisation", "34", "17%"], ["Statistics and modelling", "35", "18%"],
        ["Interpretation and report writing", "30", "15%"], ["Integration, review and reflection", "30", "15%"],
        ["Total", "195", "100%"],
    ], columns=["Work package", "Hours", "Share"]), "Effort allocation (planned)", widths_cm=[9.6, 3.5, 3.5])
    r.table(pd.DataFrame([
        ["People", "1 analyst (intern, ~32 h/week); supervisor (~1 h/week); clinical lead (2–3 h total); data protection officer (1 h)"],
        ["Computing", "Standard laptop (8–16 GB RAM); the full pipeline runs in under 15 minutes"],
        ["Software", "Open-source Python stack; Microsoft Word; optional Power BI Desktop (free)"],
        ["Data", "Public, CC BY 4.0 dataset (no cost); local data only after approval"],
        ["Budget", "No direct cost beyond staff time when open-source tools and public data are used"],
    ], columns=["Resource", "Allocation"]), "Resource plan", widths_cm=[3.0, 13.6])

    r.h("12. Risk Management and Contingency Planning", page_break=True)
    r.p("Risks are scored as likelihood × impact on 1–3 scales and reviewed at every milestone. High-impact risks "
        "have both a preventive mitigation and a contingency to apply if the risk occurs.", align="justify")
    r.figure(FIGS["risk"], "Risk heat map", None, width_cm=10.5)
    risk_table(r)
    r.h("12.1 Contingency triggers", 2)
    r.bullets(["If the dataset cannot be accessed by the end of week 1 → switch to the next-ranked source and adjust scope.",
               "If more than 40% of a key variable is missing → exclude it and record the limitation.",
               "If the model does not beat the baseline → report this honestly and focus recommendations on the statistical findings.",
               "If a required tool is unavailable → use the documented open-source alternative."])

    r.h("13. Communication Plan", page_break=True)
    r.table(pd.DataFrame([
        ["Executive leadership", "KPI headline, recommendations", "One-page brief + dashboard link", "Monthly"],
        ["Quality committee", "Findings, limitations, evidence", "Presentation + report", "Each milestone"],
        ["Clinical teams", "Risk factors, pilot design", "Ward meeting, worklist mock-up", "Weeks 3 and 5"],
        ["Supervisor", "Progress, risks, decisions", "Weekly status note", "Weekly"],
        ["Patients", "Plain-language summary", "Patient advisory group", "End of project"],
    ], columns=["Audience", "Content", "Format / channel", "Frequency"]), "Communication plan", widths_cm=[3.6, 4.6, 5.0, 3.4])

    r.h("14. Quality and Success Criteria", page_break=True)
    r.bullets(["Reproducibility: the full analysis can be regenerated from the raw download with one command.",
               "Validity: no data leakage; test data used once; assumptions checked.",
               "Transparency: every number in the report traceable to code output.",
               "Usefulness: stakeholders can name at least one action they would take from the findings.",
               "Ethics: privacy controls applied; model limitations stated; no individual decisions from the model."])
    r.h("15. Conclusion")
    r.p("This plan turns a broad aim — reduce diabetic readmissions — into a sequence of concrete, checkable steps "
        "with named owners, measurable objectives, a realistic six-week roadmap and explicit controls for the main "
        "risks. Because each step has acceptance criteria and a contingency, a reader can follow the plan without "
        "additional material.", align="justify")
    references(r, ["jencks", "raghupathi", "wirth", "kahn", "weiskopf", "meity", "hhs", "jack", "naylor", "leppin",
                   "cms", "strack", "clore", "mckinney", "pedregosa"])
    return save(r, "Week1_Strategic_Plan.docx")


# ================================================================== WEEK 2
def week2() -> Path:
    r = new_doc(2, "Data Collection and Quality Assurance Plan", "Sourcing, acquiring, validating and cleaning healthcare data")
    how_to_read(r, "This document describes how healthcare data are sourced, selected, acquired and quality-assured. "
                   "The procedures were executed on the selected dataset; Section 9 reports the actual quality "
                   "results as evidence that the plan is feasible and replicable.")
    r.h("1. Executive Summary")
    r.p(f"Seven public healthcare data sources were researched and scored on nine weighted criteria. The UCI Diabetes "
        f"130-US Hospitals dataset was selected because it is openly licensed (CC BY 4.0), immediately downloadable, "
        f"documented and contains a patient-level 30-day readmission outcome. A quality-assurance framework based on "
        f"recognised data-quality dimensions (Kahn et al., 2016; Weiskopf & Weng, 2013) defines {len(rules)} automated "
        f"validation rules, documented cleaning procedures and an issue log. Applying the plan raised cell completeness "
        f"from {qcs.loc['completeness_pct_cells', RAW]:.1f}% to {qcs.loc['completeness_pct_cells', CLEANED]:.1f}% and "
        f"all rules passed. A risk assessment identifies 12 data issues with mitigations.", align="justify")

    r.h("2. Characteristics of Healthcare Data", page_break=True)
    r.table(pd.DataFrame([
        ["Electronic health records (EHR)", "Clinical detail (diagnoses, labs, medications)", "Missingness, free text, local coding, access restrictions"],
        ["Administrative / claims data", "Large, standardised codes, costs", "Coded for billing; limited clinical detail"],
        ["Registries", "Condition-specific, curated", "Narrow scope; access agreements"],
        ["Surveys (e.g. BRFSS, NFHS)", "Population representativeness", "Self-reported; no hospital outcomes"],
        ["Hospital quality reporting (e.g. CMS)", "Comparable hospital-level measures", "Aggregated; no patient-level analysis"],
    ], columns=["Data type", "Strengths", "Typical issues"]), "Types of healthcare data", widths_cm=[4.4, 5.4, 6.8])
    r.p("Healthcare data are reused for analysis after being collected for care or payment, so their fitness for a "
        "new purpose must be demonstrated rather than assumed (Weiskopf & Weng, 2013). Data quality is also "
        "multidimensional: accuracy alone is not enough if data are incomplete, inconsistent or out of date "
        "(Wang & Strong, 1996).", align="justify")

    r.h("3. Public Data Sources Researched", page_break=True)
    r.table(pd.DataFrame([
        ["UCI Diabetes 130-US Hospitals (1999–2008)", "UCI ML Repository", "Patient encounter", "Open download, CC BY 4.0", "101,766 encounters; 30-day readmission label"],
        ["Hospital Readmissions Reduction Program", "CMS (US)", "Hospital × condition", "Open API/CSV", "FY 2026 file: 18,330 rows, 3,055 hospitals × 6 conditions"],
        ["CDC Diabetes Health Indicators (BRFSS)", "CDC via UCI", "Survey respondent", "Open download", "253,680 respondents; no hospital outcomes"],
        ["MIMIC-IV", "PhysioNet / MIT", "ICU and hospital stays", "Credentialing + training + data use agreement", "Rich clinical detail; weeks to obtain access"],
        ["Nationwide Readmissions Database", "AHRQ HCUP (US)", "Discharge record", "Purchase + data use agreement", "Designed for readmission research; cost and approval"],
        ["Global Health Observatory", "WHO", "Country indicator", "Open API", "National statistics; no patient data"],
        ["National Family Health Survey (NFHS)", "MoHFW India / DHS Program", "Household survey", "Registration request", "Indian population health; no readmission outcome"],
    ], columns=["Source", "Publisher", "Unit", "Access", "Notes"]), "Public healthcare data sources considered",
        widths_cm=[4.0, 2.8, 2.4, 3.4, 4.0], font_size=8)

    r.h("4. Selection Criteria and Process", page_break=True)
    r.table(pd.DataFrame([
        ["Relevance to the research questions", "20%", "Contains the outcome and predictors needed"],
        ["Patient-level outcome available", "15%", "Readmission flag at patient/encounter level"],
        ["Accessibility", "15%", "Can be obtained within the project timeline"],
        ["Licensing", "10%", "Reuse and redistribution permitted"],
        ["Documentation", "10%", "Variable definitions and provenance available"],
        ["Sample size", "10%", "Enough events for stable estimates and modelling"],
        ["Completeness", "10%", "Low missingness in key fields"],
        ["Recency", "5%", "Reflects current practice"],
        ["Ethical / privacy risk", "5%", "De-identified; low re-identification risk"],
    ], columns=["Criterion", "Weight", "What is assessed"]), "Weighted selection criteria", widths_cm=[5.4, 1.8, 9.4])
    r.table(pd.DataFrame([
        ["UCI Diabetes 130-US", "5", "5", "5", "5", "4", "5", "3", "1", "5", "4.50"],
        ["CMS HRRP", "4", "2", "5", "5", "4", "3", "4", "5", "5", "3.95"],
        ["CDC BRFSS indicators", "2", "1", "5", "3", "3", "5", "5", "2", "5", "3.25"],
        ["MIMIC-IV", "5", "4", "1", "2", "5", "5", "4", "4", "3", "3.70"],
        ["HCUP NRD", "5", "5", "1", "2", "5", "5", "4", "4", "4", "3.90"],
    ], columns=["Source", "Rel.", "Outcome", "Access", "Licence", "Docs", "Size", "Compl.", "Recency", "Ethics", "Weighted"]),
        "Scoring of shortlisted sources (1–5; weighted total out of 5)",
        widths_cm=[3.6, 1.1, 1.5, 1.4, 1.5, 1.2, 1.1, 1.4, 1.5, 1.3, 1.9], font_size=8,
        note="Weighted total = Σ(score × weight). Example for UCI: 5×0.20+5×0.15+5×0.15+5×0.10+4×0.10+5×0.10+3×0.10+1×0.05+5×0.05 = 4.50.")
    r.numbered(["Long-list sources from official portals (CDC, CMS, WHO, AHRQ), repositories (UCI, PhysioNet) and national surveys.",
                "Screen out sources without a usable outcome for the research questions.",
                "Verify access in practice (download or API call) and record the date.",
                "Score shortlisted sources on the weighted criteria; document the reasoning.",
                "Select the top source; keep the next two as contingencies.",
                "Do not merge sources unless populations, periods, definitions and join keys are compatible (CMS HRRP is hospital-level, 2021–2024, so it was not merged with the patient-level 1999–2008 UCI data)."])

    r.h("5. Data Acquisition Procedure", page_break=True)
    r.table(pd.DataFrame([
        ["Source", f"{src['download_url']}"], ["DOI", src["doi"]], ["Licence", src["license"]],
        ["Access date", ACCESS_DATE], ["Files", "diabetic_data.csv (encounters); IDS_mapping.csv (code labels)"],
        ["Checksum (SHA-256)", src["files"]["diabetic_data.csv"]["sha256"]],
        ["Storage", "Untouched ZIP kept in data/raw; processed outputs written to data/processed only"],
    ], columns=["Item", "Recorded value"]), "Provenance record", widths_cm=[3.6, 13.0], font_size=8)
    r.bullets(["Download with a script (no manual steps) so the process is repeatable.",
               "Extract only expected files (protects against unexpected archive contents).",
               "Record checksum, size, licence and timestamp in a metadata file.",
               "Never edit raw data; every change happens in code and is logged."])

    r.h("6. Quality Assurance Framework", page_break=True)
    r.table(pd.DataFrame([
        ["Completeness", "Are required values present?", "Missing % per column and per row; missing-pattern plot", "≤ 40% missing or documented exception"],
        ["Conformance / validity", "Do values match allowed formats and ranges?", "Range and domain rules", "0 failures"],
        ["Plausibility / consistency", "Do related fields agree?", "Cross-field rules (e.g. death vs readmission)", "< 0.5% of rows"],
        ["Uniqueness", "Are records duplicated?", "Duplicate IDs and rows; one row per patient in cohort", "0 duplicates"],
        ["Accuracy", "Are values correct?", "Comparison with published figures; clinician review", "Matches documentation"],
        ["Timeliness / relevance", "Is the data current enough?", "Data period vs intended use", "Fit for stated purpose"],
    ], columns=["Dimension", "Question", "Method", "Acceptance threshold"]), "Quality dimensions and checks (after Kahn et al., 2016)",
        widths_cm=[3.0, 4.0, 5.4, 4.2], font_size=8)
    r.h("6.1 Validation rules", 2)
    rl = rules.copy()
    rl["rule"] = rl["rule"].str.replace("'", "").str.slice(0, 70)
    r.table(rl.head(16).rename(columns={"rule": "Rule", "type": "Type", "failures": "Failures", "threshold": "Threshold", "status": "Status"}),
            f"Validation rules (first 16 of {len(rules)}; all passed after cleaning)", widths_cm=[8.0, 2.6, 1.6, 2.8, 1.6], font_size=7.5,
            note="Equivalent SQL check: SELECT COUNT(*) FROM encounters WHERE time_in_hospital NOT BETWEEN 1 AND 14; -- expected 0")

    r.h("7. Cleaning Procedures and Justification", page_break=True)
    r.table(pd.DataFrame([
        ["Placeholder codes ('?')", "Convert to true missing values", "Makes missingness countable"],
        ["Field almost entirely missing (weight, 97%)", "Drop the field", "Imputing 97% of values would fabricate data"],
        ["Missing categorical values (race, payer, specialty)", "Keep as explicit 'Unknown' level", "Avoids guessing sensitive attributes; missingness may be informative (Rubin, 1976)"],
        ["Missing numeric values", "None present; plan: multiple imputation if needed", "Preserves uncertainty (White et al., 2011)"],
        ["Invalid values (gender 'Unknown/Invalid')", "Remove the 3 rows", "Too few to form a category"],
        ["Constant columns", "Drop", "No information"],
        ["Coded IDs (admission/discharge/source)", "Map to labels with the official mapping file; group 'NULL/Not mapped' as Unknown", "Readability; consistent categories"],
        ["High-cardinality codes (700+ ICD-9)", "Group into 9 clinical categories", "Sparse levels give unstable estimates"],
        ["Repeated encounters per patient", "Keep first encounter per patient in the analysis cohort", "Independence; prevents train/test overlap"],
        ["Deaths and hospice discharges", "Exclude from the readmission cohort", "Readmission not applicable"],
    ], columns=["Issue type", "Procedure", "Justification"]), "Cleaning decisions", widths_cm=[5.0, 5.6, 6.0], font_size=8)
    tl = tlog[["step", "stage", "action", "rows_or_cells_affected"]].copy()
    tl.columns = ["#", "Stage", "Action", "Affected"]
    r.table(tl, "Transformation log (as executed)", widths_cm=[0.8, 2.6, 10.6, 2.6], font_size=7.5)

    r.h("8. Tools", page_break=True)
    r.table(pd.DataFrame([
        ["Python + pandas", "Profiling, validation and cleaning in reproducible scripts"],
        ["SQL", "Same rules expressed as queries against a hospital data warehouse"],
        ["Excel", "Quick visual inspection of samples by non-programmers (never for transformations)"],
        ["Validation libraries (e.g. Great Expectations, pandera)", "Optional: declarative rule suites for production pipelines"],
        ["Git", "Versioning of scripts and data dictionary"],
        ["Checksums (SHA-256)", "Proof that raw data have not changed"],
    ], columns=["Tool", "Role in the QA process"]), "Tools for validation, cleaning and QA", widths_cm=[5.6, 11.0])

    r.h("9. Quality Results from Executing the Plan", page_break=True)
    qshow = qc[qc["metric"].isin(["rows", "unique_patients", "completeness_pct_cells", "rows_fully_complete_pct",
                                  "columns_over_40pct_missing", "validity_pct_rows", "duplicate_encounter_ids",
                                  "constant_columns", "consistency_violations_total"])].copy()
    for c in qshow.columns[1:]:
        qshow[c] = qshow[c].map(lambda v: f"{float(v):,.2f}".rstrip("0").rstrip(".") if float(v) % 1 else f"{int(float(v)):,}")
    r.table(qshow.rename(columns={"metric": "Metric"}), "Quality metrics before and after cleaning", widths_cm=[5.6, 3.6, 3.6, 3.8], font_size=8.5,
            note="'Unknown' fills are counted as missing, so completeness gains reflect real information.")
    r.figure(F("fig01_missing_by_column.png"), "Missing data by column (raw data)", "Weight and medical specialty exceed "
             "the 40% threshold; weight was dropped and specialty kept with an explicit 'Unknown' level.", width_cm=14)
    r.figure(F("fig02_missing_pattern.png"), "Missing-data pattern (random sample)", "Gaps are not concentrated in one "
             "block of records, so they are unlikely to come from a single data-entry failure.", width_cm=14)

    r.h("10. Risk Assessment: Data Issues and Mitigation", page_break=True)
    iss = issues.rename(columns={"issue_id": "ID", "quality_dimension": "Dimension", "description": "Issue",
                                 "magnitude": "Magnitude", "resolution": "Mitigation", "residual_risk": "Residual risk"})
    r.table(iss, "Data issue risk assessment (as found)", widths_cm=[1.2, 2.0, 3.6, 2.4, 3.2, 4.2], font_size=7.5)
    r.table(pd.DataFrame([
        ["Coding changes over time (ICD-9 → ICD-10)", "Medium", "Map codes with official crosswalks; analyse eras separately"],
        ["Outcome ascertainment gaps (readmission elsewhere)", "Medium", "State as limitation; link to regional data where lawful"],
        ["Silent schema change in a refreshed extract", "Medium", "Automated schema and rule checks fail the pipeline"],
        ["Re-identification through rare combinations", "Low", "Aggregate outputs; suppress cells n < 30; no linkage"],
    ], columns=["Prospective risk", "Likelihood", "Mitigation"]), "Prospective data risks for future refreshes", widths_cm=[6.0, 2.4, 8.2], font_size=8)

    r.h("11. Governance and Privacy during Collection", page_break=True)
    r.bullets(["Use the minimum data needed; exclude identifiers from analysis.",
               "Public de-identified data: follow licence attribution (CC BY 4.0).",
               "Local data: data-protection impact assessment and approval before access; obligations under India's DPDP "
               "Act 2023 and DPDP Rules 2025 (core duties apply from 14 May 2027) and the ABDM Health Data Management "
               "Policy (MeitY, 2025; National Health Authority, 2020).",
               "US-origin data: HIPAA de-identification standards (HHS, 2012)."])
    r.h("12. Replicability Checklist")
    r.table(pd.DataFrame([["Raw data checksummed and unchanged", "✓"], ["Every transformation scripted and logged", "✓"],
                          ["Validation rules automated", "✓"], ["Data dictionary published", "✓"],
                          ["Issue log with residual risks", "✓"], ["Before/after metrics computed", "✓"]],
                         columns=["Item", "Status"]), "Replicability checklist", widths_cm=[13.6, 3.0])
    references(r, ["kahn", "weiskopf", "wang", "rubin", "white", "clore", "strack", "cms", "hhs", "meity", "nha", "mckinney"])
    return save(r, "Week2_Data_Collection_and_QA_Plan.docx")


# ================================================================== WEEK 3
def week3() -> Path:
    r = new_doc(3, "Data Visualisation and Insights Communication Plan", "Turning healthcare data into decisions for non-technical stakeholders")
    how_to_read(r, "This plan defines what to visualise, for whom, with which tools and how. Mock-ups (wireframes) show "
                   "the intended designs; the charts in Section 7 are real outputs produced when the plan was executed.")
    r.h("1. Executive Summary")
    r.p("Data only improves care when the people who make decisions understand it. This plan identifies the healthcare "
        "metrics that matter to four audiences, selects tools for each purpose, describes a repeatable process for "
        "turning raw records into accurate visuals, and provides three dashboard and report mock-ups. Design follows "
        "evidence on graphical perception (Cleveland & McGill, 1984), clarity (Tufte, 2001), narrative structure "
        "(Knaflic, 2015) and accessibility standards (W3C, 2023). Every visual is tied to a specific decision it "
        "supports.", align="justify")

    r.h("2. Audiences and Their Needs (User-Centred Design)", page_break=True)
    r.table(pd.DataFrame([
        ["Hospital executive", "Are we improving? Where should we invest?", "5 minutes, monthly", "Headline KPIs, trend vs target, top 3 messages"],
        ["Discharge planner / case manager", "Which of today's patients need extra support?", "Daily, on the ward", "Worklist with plain-language reasons and checklist"],
        ["Quality committee", "What drives readmissions? Is the evidence solid?", "Quarterly, detailed", "Rates with confidence intervals, adjusted results, data-quality status"],
        ["Clinician (diabetes team)", "Is HbA1c testing consistent?", "Monthly", "Testing rate by diagnosis and ward"],
    ], columns=["Persona", "Key question", "Context", "What they need to see"]), "Audience personas", widths_cm=[3.6, 4.6, 3.0, 5.4], font_size=8.5)
    r.p("Design principle: start from the decision, not the data. Each persona's key question defines the first "
        "thing they see; detail is available on demand (filters, tooltips, data tables).", align="justify")

    r.h("3. Key Metrics to Visualise", page_break=True)
    r.table(pd.DataFrame([
        ["30-day readmission rate", "Core quality outcome", "KPI card with 95% CI; line chart over time", "All"],
        ["Rate by discharge destination", "Identifies where transition support is needed", "Sorted horizontal bars with CIs", "Quality, case managers"],
        ["Rate by prior admissions", "Simple risk-stratification signal", "Ordered bars (0, 1, 2, 3+)", "Case managers"],
        ["Rate by age and diagnosis", "Targets clinical pathways", "Bars / small multiples", "Quality committee"],
        ["Length of stay distribution", "Resource use, outliers", "Histogram; box plot by diagnosis", "Executives, bed managers"],
        ["Bed-day share by diagnosis", "Where capacity goes", "Sorted bars", "Executives"],
        ["HbA1c testing rate", "Care-process gap", "Bars by diagnosis; trend", "Diabetes team"],
        ["Data completeness", "Trust in the numbers", "Traffic-light list with text labels", "Quality, HIM"],
        ["Model risk deciles", "Whether a risk score separates patients", "Bars of observed rate by decile", "Quality, analytics"],
    ], columns=["Metric", "Why it matters", "Recommended visual", "Audience"]), "Metrics and visual encodings", widths_cm=[4.0, 4.4, 5.0, 3.2], font_size=8)

    r.h("4. Tool Selection and Justification", page_break=True)
    r.table(pd.DataFrame([
        ["Power BI", "Interactive dashboards, DAX measures, row-level security, widely used in hospitals", "Licence for sharing; Windows desktop authoring", "Primary production dashboard"],
        ["Tableau", "Strong visual exploration", "Licence cost", "Alternative if already licensed"],
        ["Python (Matplotlib, seaborn)", "Exact control, reproducible, publication quality", "Static; needs coding", "Report figures"],
        ["HTML + JavaScript", "Runs offline in any browser; no licence", "Development effort", "Portable dashboard (delivered)"],
        ["Excel", "Familiar to all staff", "Error-prone for complex logic; weak interactivity", "Ad-hoc tables only"],
    ], columns=["Tool", "Strengths", "Limitations", "Role in this plan"]), "Tool comparison", widths_cm=[3.0, 5.4, 4.4, 3.8], font_size=8)
    r.p("Choice: Python for the analytical figures (reproducible from code), a self-contained HTML dashboard for "
        "immediate use without licences, and a Power BI build pack (star-schema data model and DAX measures) for "
        "organisations that standardise on Power BI.", align="justify")

    r.h("5. From Raw Data to Visual: Process", page_break=True)
    r.numbered(["Start from the quality-assured analysis cohort (never from raw data).",
                "Aggregate to the level of the decision (e.g. rate per discharge group).",
                "Compute uncertainty: 95% Wilson confidence intervals for every rate (Wilson, 1927).",
                "Suppress groups with fewer than 30 patients (privacy and stability).",
                "Choose the encoding by the question: position and length are read most accurately (Cleveland & McGill, 1984).",
                "Apply the design system: one colour per series, colour-blind-safe palette, no dual axes.",
                "Write a title that states the finding and an interpretation sentence.",
                "Validate: check every number against the source table; test with one target user.",
                "Publish with method notes and a data-table alternative."])
    r.h("5.1 Design and accessibility standards", 2)
    r.table(pd.DataFrame([
        ["Colour", "Colour-vision-deficiency-safe palette; never rely on colour alone (add labels or icons)"],
        ["Contrast", "Text contrast at least 4.5:1 (WCAG 2.2)"],
        ["Alternatives", "Every chart has a data table and alt text"],
        ["Keyboard and screen readers", "Filters usable by keyboard; charts labelled for assistive technology"],
        ["Language", "Plain-language titles; avoid jargon such as 'AUC' on executive views"],
        ["Uncertainty", "Show confidence intervals; grey out small groups"],
        ["Integrity", "Zero-based bars; no 3-D; no dual axes; consistent scales between periods"],
    ], columns=["Aspect", "Standard"]), "Design and accessibility standards", widths_cm=[3.6, 13.0], font_size=8.5)

    r.h("6. Mock-ups and Wireframes", page_break=True)
    r.figure(FIGS["wf_exec"], "Wireframe: executive overview dashboard", "Answers 'are we improving and where should we act?' "
             "in under five minutes: KPI cards on top, trend against target, the main driver (discharge destination) "
             "and plain-language key messages.")
    r.figure(FIGS["wf_ops"], "Wireframe: discharge-planning worklist", "Turns analysis into daily action: patients ranked "
             "into priority groups, with reasons in plain language and a discharge checklist. Clinician judgement always "
             "overrides the flag.")
    r.figure(FIGS["wf_report"], "Wireframe: quarterly one-page quality brief", "A printable page with a headline "
             "sentence, two consistent charts, a data-quality traffic light (with text labels) and owned actions.")
    r.h("6.1 Implemented dashboard", 2)
    r.figure(F("dashboard_screenshot_desktop.png"), "Implemented interactive dashboard (desktop view)", "The executed version "
             "includes five filters, six KPI cards with confidence intervals, six charts with tooltips and data tables, "
             "and automatic suppression of small groups.", width_cm=12)

    r.h("7. Visualisations and the Decisions They Support", page_break=True)
    for fig, cap, decision in [
        ("fig00_executive_kpis.png", "Executive KPI overview", f"Sets the baseline: {pct(ov['rate'])} of patients readmitted within 30 days; supports target setting."),
        ("fig08_readmission_by_discharge.png", "Readmission by discharge destination", f"Problem identified: facility discharges are readmitted at {pct(fac['rate_1'])} vs {pct(fac['rate_0'])}. Decision: prioritise facility hand-offs."),
        ("fig07_prior_inpatient.png", "Readmission by prior admissions", "Problem identified: a steep risk gradient. Decision: flag patients with prior admissions for early discharge planning."),
        ("fig09_hba1c.png", "HbA1c testing and readmission", f"Problem identified: only {pct(eda['a1c_measured_rate'])} tested. Decision: audit and pilot consistent testing."),
        ("fig11_utilisation_by_diagnosis.png", "Bed-day consumption by diagnosis", "Supports capacity planning: circulatory conditions use the most bed-days."),
        ("fig17_feature_importance.png", "What the risk model relies on", "Builds trust: clinicians can see the model uses clinically sensible signals."),
    ]:
        r.figure(F(fig), cap, decision, width_cm=14)

    r.h("8. Communicating with Non-Technical Stakeholders", page_break=True)
    r.bullets([("Lead with the answer: ", "a headline sentence before any chart (e.g. 'One in eleven patients returns within 30 days')."),
               ("Use a story arc: ", "context → problem → evidence → recommended action → expected benefit (Knaflic, 2015)."),
               ("Translate statistics: ", "'about twice the risk' rather than 'RR 2.03 (95% CI 1.93–2.13)'; keep exact numbers in the appendix."),
               ("Show uncertainty simply: ", "error bars plus a sentence such as 'the true value is very likely between x and y'."),
               ("Be explicit about limits: ", "associations, not causes; old data; model not for individual decisions."),
               ("Close the loop: ", "every dashboard page names an owner and an action.")])
    r.table(pd.DataFrame([
        ["Executive brief", "One page + dashboard link", "Monthly"], ["Quality committee pack", "Report + live dashboard walk-through", "Quarterly"],
        ["Ward huddle", "Worklist + 3 key messages", "Weekly"], ["Patient advisory group", "Plain-language summary", "Twice yearly"],
    ], columns=["Product", "Format", "Cadence"]), "Communication products", widths_cm=[5.0, 7.6, 4.0])
    r.h("9. Evaluating the Visualisations")
    r.bullets(["Usability test with one user per persona: can they state the main finding in 30 seconds?",
               "Accuracy check: every displayed number reconciles with the analysis tables.",
               "Accessibility check: contrast, keyboard use, screen-reader labels, colour-blind simulation.",
               "Adoption metrics: dashboard visits, actions logged from the worklist, feedback scores."])
    references(r, ["cleveland", "tufte", "knaflic", "wcag", "wilson", "clore", "strack"])
    return save(r, "Week3_Data_Visualisation_Plan.docx")


# ================================================================== WEEK 4
def week4() -> Path:
    r = new_doc(4, "Statistical Analysis and Predictive Modelling Framework", "Methods, validation and ethics for healthcare predictive analytics")
    how_to_read(r, "This framework is written so that it can be replicated on another hospital's data. Section 10 shows "
                   "the results obtained when the framework was applied to the selected dataset, demonstrating that it "
                   "is practical.")
    r.h("1. Executive Summary")
    r.p("The framework moves from descriptive statistics to inferential tests, adjusted regression and finally "
        "predictive modelling, with each method chosen for the data type and question. Data are segmented by patient "
        "into training, validation and test sets; variables are selected on clinical grounds and screened for "
        "leakage; models are compared against simple baselines and validated by cross-validation and a single "
        "untouched test evaluation, including calibration and subgroup analysis. Ethical safeguards cover privacy, "
        "fairness, transparency, accountability and human oversight, consistent with TRIPOD+AI reporting "
        "(Collins et al., 2024), WHO (2021) and ICMR (2023) guidance.", align="justify")

    r.h("2. Statistical Techniques for Healthcare Data", page_break=True)
    r.table(pd.DataFrame([
        ["Descriptive statistics", "Summarise distributions (mean, median, IQR, rates)", "Use medians for skewed counts"],
        ["Confidence intervals (Wilson)", "Precision of rates", "Binomial sampling"],
        ["Chi-square test", "Categorical factor vs binary outcome", "Independent observations; expected counts ≥ 5"],
        ["Mann–Whitney U / Kruskal–Wallis", "Compare skewed numeric variables between groups", "Independent groups"],
        ["t-test / ANOVA", "Compare means of roughly normal variables", "Normality, equal variances (or Welch)"],
        ["Correlation (Spearman)", "Monotonic association", "No causal meaning"],
        ["Logistic regression", "Adjusted associations with a binary outcome", "Linear logit; low collinearity"],
        ["Poisson / negative binomial regression", "Counts such as length of stay or visits", "Over-dispersion check"],
        ["Survival analysis (Kaplan–Meier, Cox)", "Time to readmission with censoring", "Proportional hazards"],
        ["Time-series (ETS, ARIMA)", "Trends and forecasts of monthly rates", "Requires dated data"],
        ["Multiple-comparison control (Holm)", "Many tests at once", "Controls family-wise error (Holm, 1979)"],
    ], columns=["Technique", "Use in healthcare", "Key assumption / note"]), "Common statistical techniques", widths_cm=[4.6, 6.0, 6.0], font_size=8)

    r.h("3. Mapping Questions to Methods", page_break=True)
    r.table(pd.DataFrame([
        ["How common is readmission?", "Proportion + Wilson CI", "Rate with 95% CI"],
        ["Which factors are associated?", "χ², Mann–Whitney (+ effect sizes, Holm)", "Cramér's V, rank-biserial r"],
        ["Which associations survive adjustment?", "Multivariable logistic regression", "Adjusted odds ratios"],
        ["Does an association vary by subgroup?", "Interaction terms + likelihood-ratio test", "Effect modification"],
        ["How are resources used?", "Kruskal–Wallis; bed-day shares", "ε², distributions"],
        ["Can risk be predicted?", "Supervised machine learning", "ROC-AUC, PR-AUC, calibration"],
        ["How will rates change? (when dates exist)", "Exponential smoothing / ARIMA", "Forecast with prediction intervals"],
    ], columns=["Question", "Method", "Main output"]), "Question-to-method mapping", widths_cm=[5.4, 6.0, 5.2])

    r.h("4. Data Segmentation", page_break=True)
    r.bullets([("Cohort definition: ", "one index encounter per patient; exclude deaths and hospice discharges (outcome not applicable)."),
               ("Unit of analysis: ", "the patient — prevents the same person appearing in training and test data."),
               ("Split: ", "60% training, 20% validation, 20% test, stratified on the outcome so each set has the same readmission rate."),
               ("Grouped or temporal splits: ", "if repeated encounters are used, split by patient (GroupKFold); if dates exist, train on earlier periods and test on later ones."),
               ("Analytic segments: ", "pre-specified subgroups (age, sex, race, diagnosis) for descriptive comparison and fairness audit.")])
    r.figure(FIGS["modelling"], "Step-by-step predictive modelling process", None)

    r.h("5. Variable Selection", page_break=True)
    r.numbered(["Fix the prediction time (at discharge) and keep only variables known by then.",
                "Start from clinical knowledge and literature (prior use, comorbidity, discharge destination, age).",
                "Remove identifiers and anything derived from the outcome (target leakage; Kaufman et al., 2012).",
                "Group sparse categories; check collinearity (variance inflation factor < 5).",
                "Decide on sensitive attributes deliberately: race and payer are excluded as predictors but used to audit fairness.",
                "Let regularisation (L2) and tree models handle remaining redundancy; confirm with permutation importance."])
    r.table(pd.DataFrame(list(mdl["excluded_features"].items()), columns=["Excluded variable(s)", "Reason"]),
            "Variables excluded and why", widths_cm=[6.0, 10.6], font_size=8.5)

    r.h("6. Pre-processing and Model Selection", page_break=True)
    r.table(pd.DataFrame([
        ["Missing values", "Explicit 'Unknown' level for categories; multiple imputation for numeric values if needed (White et al., 2011)"],
        ["Encoding", "One-hot encoding; rare levels grouped"],
        ["Scaling", "Standardisation for linear models only"],
        ["Class imbalance", "Keep natural prevalence (preserves calibration); handle through threshold choice"],
        ["Leakage prevention", "All steps fitted inside a pipeline on training folds only"],
    ], columns=["Step", "Approach"]), "Pre-processing", widths_cm=[4.0, 12.6])
    r.table(pd.DataFrame([
        ["Baseline: prevalence", "Reference for no information", "AUC = 0.5 by definition"],
        ["Baseline: one-variable rule", "What a hospital could do today", "Prior-admission flag"],
        ["Logistic regression", "Interpretable, well calibrated", "Linear effects only"],
        ["Random forest", "Non-linearities and interactions", "Less transparent; can be poorly calibrated"],
        ["Gradient boosting", "Strong tabular performance", "Needs careful tuning"],
    ], columns=["Model", "Why", "Trade-off"]), "Candidate models", widths_cm=[4.6, 6.0, 6.0])
    r.p("Hyperparameters are tuned by randomised search with 5-fold stratified cross-validation on the training set. "
        "The final model is chosen on the validation set, with a parsimony rule favouring logistic regression when "
        "performance is within 0.005 AUC (Steyerberg, 2019; Hastie et al., 2009).", align="justify")

    r.h("7. Evaluation and Validation", page_break=True)
    r.table(pd.DataFrame([
        ["k-fold cross-validation", "Tuning and internal performance estimate", "Training set only"],
        ["Hold-out validation set", "Model and threshold selection", "Never used for training"],
        ["Untouched test set", "Final, unbiased performance", "Used exactly once"],
        ["Bootstrap", "Confidence intervals for AUC and PR-AUC", "1,000 resamples"],
        ["Temporal / external validation", "Transportability to new periods or hospitals", "Required before deployment"],
    ], columns=["Technique", "Purpose", "Rule"]), "Validation design", widths_cm=[4.6, 7.0, 5.0])
    r.table(pd.DataFrame([
        ["Discrimination", "ROC-AUC; PR-AUC (better for rare outcomes; Saito & Rehmsmeier, 2015)"],
        ["Calibration", "Calibration curve, slope and intercept (Van Calster et al., 2019)"],
        ["Classification at a threshold", "Sensitivity, specificity, PPV, NPV, F1 — threshold chosen from workload"],
        ["Clinical utility", "Decision-curve analysis (net benefit; Vickers & Elkin, 2006)"],
        ["Fairness", "Metrics by subgroup with confidence intervals"],
    ], columns=["Aspect", "Metrics"]), "Evaluation metrics", widths_cm=[4.0, 12.6])
    r.p(f"Accuracy is misleading for rare outcomes: in this data a rule that flags nobody is "
        f"{pct(thr.iloc[-1]['accuracy'])} accurate but finds no readmissions.", align="justify")

    r.h("8. Pitfalls, Biases and Mitigation", page_break=True)
    r.table(pd.DataFrame([
        ["Target leakage", "Inflated performance", "Prediction-time rule; exclude outcome-derived fields; pipelines"],
        ["Non-independent records", "Over-optimistic validation", "Patient-level split"],
        ["Overfitting", "Poor generalisation", "Cross-validation; regularisation; simple baselines"],
        ["Confounding", "Spurious causal claims", "Adjusted models; causal language avoided"],
        ["Selection bias", "Unrepresentative cohort", "Document inclusion rules; compare with population"],
        ["Measurement / label bias", "Model learns documentation patterns", "Audit missingness by group; clinician review"],
        ["Historical bias", "Model reproduces inequities (Obermeyer et al., 2019)", "Exclude proxies; fairness audit"],
        ["Dataset shift", "Performance decays over time", "Temporal validation; drift monitoring"],
        ["Multiple testing", "False positives", "Holm correction; pre-specified hypotheses"],
    ], columns=["Pitfall", "Consequence", "Mitigation"]), "Pitfalls and mitigation", widths_cm=[3.8, 5.0, 7.8], font_size=8.5)

    r.h("9. Ethical Considerations", page_break=True)
    r.table(pd.DataFrame([
        ["Privacy", "De-identified data; minimum necessary; secure storage; DPDP Act/Rules obligations for local data (MeitY, 2025)"],
        ["Fairness and equity", "Exclude race and payer as inputs; report performance by subgroup; agree fairness criteria with stakeholders (Rajkomar et al., 2018)"],
        ["Transparency", "Model card; plain-language reasons; published code"],
        ["Accountability", "Named clinical owner; governance committee sign-off"],
        ["Autonomy and human oversight", "Scores support, never replace, clinical judgement"],
        ["Beneficence / non-maleficence", "Weigh harms of false negatives (missed patients) and false positives (unnecessary follow-up)"],
        ["Validity before use", "External validation and recalibration before deployment (ICMR, 2023; WHO, 2021)"],
        ["Monitoring", "Quarterly drift, calibration and subgroup checks; decommissioning plan"],
    ], columns=["Principle", "How it is applied"]), "Ethical framework", widths_cm=[4.0, 12.6], font_size=8.5)

    r.h("10. Illustrative Application of the Framework", page_break=True)
    r.p(f"Applied to {num(ov['n'])} patients, the framework produced the following results.", align="justify")
    r.bullets([f"Facility discharge: {pct(fac['rate_1'])} vs {pct(fac['rate_0'])} readmitted (risk ratio {fac['risk_ratio']:.2f}); adjusted OR {OR_REHAB:.2f} (rehab) and {OR_SNF:.2f} (skilled nursing) vs home.",
               f"Each prior inpatient admission: adjusted OR {OR_IP:.2f}.",
               f"HbA1c measured: adjusted OR {inter['a1c_measured_OR_main_model']:.2f}; association varies by diagnosis (p = {pval(inter['p_value'])}).",
               f"Best model ({best}): test ROC-AUC {mdl['test_roc_auc']:.3f} (95% CI {mdl['test_roc_auc_ci'][0]:.3f}–{mdl['test_roc_auc_ci'][1]:.3f}) vs {base_ip['roc_auc']:.3f} for the one-variable rule; calibration slope {cal['calibration_slope']:.2f}."])
    ts = test[["model", "roc_auc", "pr_auc", "brier", "recall_sensitivity", "precision_ppv"]].copy()
    for c in ts.columns[1:]:
        ts[c] = ts[c].map(lambda v: f"{v:.3f}")
    ts.columns = ["Model", "ROC-AUC", "PR-AUC", "Brier", "Sensitivity", "Precision"]
    r.table(ts, "Test-set results", widths_cm=[5.8, 2.1, 2.1, 2.0, 2.4, 2.2], font_size=8)
    r.figure(F("fig13_logistic_regression_forest.png"), "Adjusted odds ratios (logistic regression)", None, width_cm=14)
    r.figure(F("fig14_roc_pr_curves.png"), "ROC and precision–recall curves", None, width_cm=15)
    r.figure(F("fig15_calibration.png"), "Calibration curves", None, width_cm=9.5)
    r.figure(F("fig18_subgroup_performance.png"), "Subgroup performance (fairness audit)", None, width_cm=15)
    r.h("11. Replication Checklist")
    r.table(pd.DataFrame([["Prediction time and unit defined", "✓"], ["Leakage screen documented", "✓"],
                          ["Patient-level split, fixed seed", "✓"], ["Pre-processing inside pipelines", "✓"],
                          ["Baselines reported", "✓"], ["Test set used once", "✓"], ["Calibration and fairness reported", "✓"],
                          ["External / temporal validation", "Required before use"]],
                         columns=["Item", "Status"]), "Replication checklist", widths_cm=[11.6, 5.0])
    references(r, ["tripod", "who", "icmr", "holm", "kaufman", "white", "steyerberg", "hastie", "saito", "vancalster",
                   "vickers", "obermeyer", "rajkomar", "meity", "wilson", "kansagara"])
    return save(r, "Week4_Statistical_Analysis_and_Modelling_Framework.docx")


# ================================================================== WEEK 5
def week5() -> Path:
    r = new_doc(5, "Analytical Report and Evidence-Based Recommendations", "Findings from a diabetic inpatient readmission analysis")
    r.h("1. Executive Summary")
    r.p(f"We analysed {num(ov['n'])} patients with diabetes admitted to 130 US hospitals (first eligible stay each). "
        f"{pct(ov['rate'])} were readmitted within 30 days — roughly one patient in eleven.", align="justify")
    r.p("Key insights", bold_lead="")
    r.doc.paragraphs[-1].runs[0].bold = True
    r.bullets([f"Patients discharged to a care facility return about twice as often ({pct(fac['rate_1'])} vs {pct(fac['rate_0'])}).",
               f"Patients admitted in the previous year return about twice as often ({pct(ip['rate_1'])} vs {pct(ip['rate_0'])}).",
               f"HbA1c was checked for only {pct(eda['a1c_measured_rate'])} of patients; testing is linked to slightly fewer readmissions.",
               f"Average stay is {eda['mean_los']:.1f} days; circulatory conditions use {pct(util.iloc[0]['bed_day_share'])} of bed-days.",
               f"A risk model can rank patients (the top 20% contain {pct(m20['recall_sensitivity'])} of readmissions) but is not accurate enough for individual decisions."])
    r.p("Top recommendations", bold_lead="")
    r.doc.paragraphs[-1].runs[0].bold = True
    r.numbered(["Strengthen hand-offs for patients discharged to care facilities.",
                "Start discharge planning early and book follow-up within 7 days for patients with a prior-year admission.",
                "Audit HbA1c testing and pilot a prompt with a controlled evaluation.",
                "Improve recording of weight, specialty and payer.",
                "Use any risk model only in a supervised, locally validated pilot."])

    r.h("2. Introduction", page_break=True)
    r.p("Readmissions within 30 days are costly, often avoidable and publicly reported (Jencks et al., 2009; CMS, 2026). "
        "This report answers four questions for a hospital quality team: how common early readmission is among "
        "diabetic inpatients; which factors are associated with it; where inpatient resources are used; and whether "
        "readmission risk can be predicted at discharge.", align="justify")

    r.h("3. Methodology", page_break=True)
    r.table(pd.DataFrame([
        ["Data", f"UCI Diabetes 130-US Hospitals (1999–2008), {num(clean['raw_rows'])} encounters, CC BY 4.0 (Clore et al., 2014)"],
        ["Cohort", f"First eligible encounter per patient; deaths and hospice excluded → {num(ov['n'])} patients"],
        ["Quality assurance", f"{len(rules)} validation rules (all passed); documented cleaning; issue log"],
        ["Outcome", "Readmission within 30 days of discharge"],
        ["Statistics", "Rates with 95% Wilson CIs; χ² and Mann–Whitney tests with effect sizes and Holm correction; multivariable logistic regression"],
        ["Prediction", "Logistic regression, random forest, gradient boosting vs baselines; patient-level 60/20/20 split; test set used once"],
        ["Tools", "Python (pandas, SciPy, statsmodels, scikit-learn), HTML dashboard, Power BI build pack"],
    ], columns=["Element", "Approach"]), "Methodology summary", widths_cm=[3.4, 13.2], font_size=8.5)
    r.p("How insights were derived: every finding below comes from a computed table, is shown with its uncertainty, "
        "and was checked against an adjusted model before being used for a recommendation. Effect sizes, not just "
        "p-values, determine importance, because with 70,000 patients even trivial differences are 'significant'.",
        align="justify")

    r.h("4. Results", page_break=True)
    r.figure(F("fig00_executive_kpis.png"), "Headline indicators", None)
    r.h("4.1 Who is readmitted?", 2)
    r.figure(F("fig08_readmission_by_discharge.png"), "Readmission by discharge destination",
             f"Facility discharges carry the highest risk ({pct(dis.loc['Other inpatient / rehab facility', 'rate'])} rehab/other inpatient, "
             f"{pct(dis.loc['Skilled nursing / ICF', 'rate'])} skilled nursing) versus {pct(dis.loc['Home', 'rate'])} for home.", width_cm=14)
    r.figure(F("fig07_prior_inpatient.png"), "Readmission by prior-year admissions",
             f"Risk rises from {pct(by_ip.iloc[0]['rate'])} (none) to {pct(by_ip.iloc[-1]['rate'])} (five or more).", width_cm=14)
    r.figure(F("fig04_age_distribution_readmission.png"), "Age profile and readmission",
             f"Patients aged 60+ are readmitted more often ({pct(old['rate_1'])} vs {pct(old['rate_0'])}).", width_cm=15)
    r.h("4.2 Care processes", 2)
    r.figure(F("fig09_hba1c.png"), "HbA1c testing and readmission",
             f"Adjusted OR for testing {inter['a1c_measured_OR_main_model']:.2f} (95% CI {inter['a1c_measured_OR_ci'][0]:.2f}–{inter['a1c_measured_OR_ci'][1]:.2f}); varies by diagnosis.", width_cm=15)
    r.h("4.3 Resource use", 2)
    r.figure(F("fig11_utilisation_by_diagnosis.png"), "Bed-days and length of stay by diagnosis",
             f"Stays of 7+ days are linked to more readmissions ({pct(los7['rate_1'])} vs {pct(los7['rate_0'])}).", width_cm=15)
    r.h("4.4 Adjusted analysis", 2)
    r.figure(F("fig13_logistic_regression_forest.png"), "Adjusted odds ratios",
             f"Discharge destination (OR up to {OR_REHAB:.2f}), prior admissions (OR {OR_IP:.2f} each) and age remain the strongest associations after adjustment.", width_cm=14)
    r.h("4.5 Predicting readmission", 2)
    dshow = dec[["decile", "patients", "readmissions", "observed_rate", "capture_cumulative"]].copy()
    dshow["observed_rate"] = dshow["observed_rate"].map(pct)
    dshow["capture_cumulative"] = dshow["capture_cumulative"].map(pct)
    dshow.columns = ["Risk decile (1 = highest)", "Patients", "Readmissions", "Observed rate", "Cumulative share of readmissions"]
    r.table(dshow, "Observed readmission by predicted-risk decile (test set)", widths_cm=[3.6, 2.4, 2.6, 3.0, 5.0], font_size=8)
    r.p(f"The highest-risk tenth of patients is readmitted at {pct(mdl['top_decile_rate'])}, five times the lowest tenth "
        f"({pct(mdl['bottom_decile_rate'])}). Overall discrimination is modest (ROC-AUC {mdl['test_roc_auc']:.3f}), typical "
        f"of administrative-data models (Kansagara et al., 2011).", align="justify")

    r.h("5. Discussion", page_break=True)
    r.p("Readmission risk concentrates in patients who are frailer and more dependent at discharge. These markers are "
        "mostly not levers a hospital can change directly, so their value is in targeting: directing follow-up to the "
        "groups that carry most of the risk. The HbA1c finding points to a care-process gap that is worth testing.", align="justify")
    r.bullets(["Alternative explanations: facility discharge and prior admissions may stand in for unmeasured illness severity or social support.",
               "HbA1c testing may mark more attentive care rather than cause better outcomes.",
               "Readmissions to other hospitals may be missed, biasing rates downward."])

    r.h("6. Evidence-Based Recommendations", page_break=True)
    rec = pd.DataFrame([
        ["Facility discharge = highest risk", f"{pct(fac['rate_1'])} vs {pct(fac['rate_0'])}; adj. OR {OR_REHAB:.2f}", "Medication list and summary sent before transfer; 72-h call to facility", "Case management", "High"],
        ["Prior admissions = steep gradient", f"{pct(ip['rate_1'])} vs {pct(ip['rate_0'])}; adj. OR {OR_IP:.2f}", "Flag at admission; follow-up ≤ 7 days", "Hospitalists", "High"],
        ["HbA1c rarely tested", f"{pct(eda['a1c_measured_rate'])} tested; adj. OR {inter['a1c_measured_OR_main_model']:.2f}", "Audit; pilot testing prompt with evaluation", "Diabetes team", "Medium"],
        ["Older patients higher risk", f"Age ≥ 60 RR {old['risk_ratio']:.2f}", "Functional/caregiver assessment for ≥ 70 (to be tested)", "Geriatrics", "Medium"],
        ["Documentation gaps", f"Weight {miss['weight']:.0f}%, specialty {miss['medical_specialty']:.0f}% missing", "Mandatory structured fields; quality dashboard", "HIM / IT", "Medium"],
        ["Model modest", f"AUC {mdl['test_roc_auc']:.3f}", "Supervised pilot only after local validation", "Clinical governance", "Medium"],
    ], columns=["Finding", "Evidence", "Action", "Owner", "Priority"])
    r.table(rec, "Recommendations traced to evidence", widths_cm=[3.4, 3.6, 5.0, 2.6, 2.0], font_size=8)
    r.p("Supporting external evidence: structured discharge programmes reduced 30-day readmissions in randomised trials "
        "(Jack et al., 2009; Naylor et al., 1999; Leppin et al., 2014).", align="justify")

    r.h("7. Communicating the Findings", page_break=True)
    r.p("Plain-language summary for patients and non-specialists:", bold_lead="")
    r.doc.paragraphs[-1].runs[0].bold = True
    r.p("About one in eleven patients with diabetes came back to hospital within a month of going home. People going to a "
        "nursing or rehabilitation facility, and people who had been in hospital before, came back about twice as often. "
        "Hospitals can use this to give these patients extra help when they leave — for example making sure the facility "
        "has the right medicines list and that a follow-up visit is booked.", italic=True, align="justify")
    r.bullets(["Executives: one-page brief + KPI dashboard (Figure 1).",
               "Case managers: discharge worklist with plain-language reasons.",
               "Quality committee: this report plus the interactive dashboard."])

    r.h("8. Limitations and Future Research", page_break=True)
    r.bullets(["Data are from 1999–2008 (ICD-9); rates may not reflect current care.",
               "No dates or hospital IDs: no trends or hospital comparisons.",
               "Observational data: associations, not causes.",
               "Model not externally validated; calibration needs adjustment.",
               "Future work: current local data, time-to-event models, decision-curve analysis, and controlled evaluation of the recommended interventions."])
    r.h("9. Conclusion")
    r.p("The analysis identifies clear, actionable concentrations of readmission risk and a care-process gap. Targeted "
        "transition support for facility discharges and frequent users, better diabetes assessment and better data "
        "capture are the evidence-based priorities.", align="justify")
    references(r, ["jencks", "cms", "clore", "strack", "kansagara", "jack", "naylor", "leppin"])
    return save(r, "Week5_Analytical_Report_and_Recommendations.docx")


def week6(final_report: Path) -> Path:
    out = OUT_DIR / "Week6_Integrated_Project_Report.docx"
    shutil.copy(final_report, out)
    print(f"Saved {out.name} (integrated report from step 7)")
    return out


def main() -> None:
    week1(); week2(); week3(); week4(); week5()
    week6(ROOT / "report" / "Healthcare_Data_Analytics_Final_Report.docx")


if __name__ == "__main__":
    main()
