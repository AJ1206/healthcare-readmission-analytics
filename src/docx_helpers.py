"""python-docx helpers: page setup, fields (TOC, captions, page numbers), styled tables and lists."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

ACCENT = RGBColor(0x1C, 0x5C, 0xAB)
INK = RGBColor(0x1F, 0x1F, 0x1F)
MUTED = RGBColor(0x59, 0x59, 0x59)
HEADER_FILL = "1C5CAB"
BAND_FILL = "F2F6FB"
MARGIN = Cm(2.2)


class Report:
    def __init__(self) -> None:
        self.doc = Document()
        self.fig_n = 0
        self.tab_n = 0
        self._setup()

    # ------------------------------------------------------------------ setup
    def _setup(self) -> None:
        s = self.doc.sections[0]
        s.page_width, s.page_height = Mm(210), Mm(297)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(s, side, MARGIN)
        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = "Calibri"
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = INK
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        for lvl, size, before in [(1, 16, 18), (2, 13, 12), (3, 11.5, 10)]:
            st = styles[f"Heading {lvl}"]
            st.font.name = "Calibri"
            rfonts = st.element.rPr.rFonts
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                rfonts.attrib.pop(qn(attr), None)  # theme fonts would override the explicit font
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.color.rgb = ACCENT if lvl < 3 else INK
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.keep_with_next = True
        cap = styles["Caption"]
        cap.font.name = "Calibri"
        cap.font.size = Pt(9)
        cap.font.italic = False
        cap.font.bold = False
        cap.font.color.rgb = MUTED
        cap.paragraph_format.space_after = Pt(10)
        for name in ("List Bullet", "List Number"):
            styles[name].font.size = Pt(10.5)
            styles[name].paragraph_format.space_after = Pt(3)

    def page_numbers(self, header_text: str) -> None:
        s = self.doc.sections[0]
        s.different_first_page_header_footer = True
        hp = s.header.paragraphs[0]
        hp.text = header_text
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r in hp.runs:
            r.font.size = Pt(8)
            r.font.color.rgb = MUTED
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._run(fp, "Page ", 8.5)
        self._field(fp, "PAGE", "1", 8.5)
        self._run(fp, " of ", 8.5)
        self._field(fp, "NUMPAGES", "1", 8.5)

    @staticmethod
    def _run(p, text, size=None, bold=False, color=None, italic=False):
        r = p.add_run(text)
        if size:
            r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        if color is not None:
            r.font.color.rgb = color
        return r

    @staticmethod
    def _field(paragraph, instr: str, placeholder: str = "", size=None):
        def fld(tag, **attrs):
            el = OxmlElement(tag)
            for k, v in attrs.items():
                el.set(qn(k), v)
            return el
        r1 = paragraph.add_run()
        r1._r.append(fld("w:fldChar", **{"w:fldCharType": "begin", "w:dirty": "true"}))
        r2 = paragraph.add_run()
        it = fld("w:instrText", **{"xml:space": "preserve"})
        it.text = f" {instr} "
        r2._r.append(it)
        r3 = paragraph.add_run()
        r3._r.append(fld("w:fldChar", **{"w:fldCharType": "separate"}))
        r4 = paragraph.add_run(placeholder)
        r5 = paragraph.add_run()
        r5._r.append(fld("w:fldChar", **{"w:fldCharType": "end"}))
        if size:
            for r in (r1, r2, r3, r4, r5):
                r.font.size = Pt(size)
        return r4

    # ------------------------------------------------------------------ blocks
    def h(self, text: str, level: int = 1, page_break: bool = False):
        if page_break:
            self.page_break()
        return self.doc.add_heading(text, level=level)

    def p(self, text: str = "", bold_lead: str | None = None, size=None, italic=False, align=None, color=None,
          space_after=None):
        para = self.doc.add_paragraph()
        if bold_lead:
            self._run(para, bold_lead, size, bold=True)
        self._run(para, text, size, italic=italic, color=color)
        if align == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "justify":
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if space_after is not None:
            para.paragraph_format.space_after = Pt(space_after)
        return para

    def bullets(self, items, style: str = "List Bullet"):
        for it in items:
            para = self.doc.add_paragraph(style=style)
            if isinstance(it, tuple):
                self._run(para, it[0], bold=True)
                self._run(para, it[1])
            else:
                self._run(para, it)

    def numbered(self, items):
        """Numbered list that restarts at 1 for every call (Word's List Number style continues across lists)."""
        for i, it in enumerate(items, start=1):
            para = self.doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.75)
            para.paragraph_format.first_line_indent = Cm(-0.5)
            para.paragraph_format.space_after = Pt(3)
            self._run(para, f"{i}.  ")
            self._run(para, it)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def caption(self, kind: str, text: str):
        para = self.doc.add_paragraph(style="Caption")
        self._run(para, f"{kind} ", bold=True)
        n = self.fig_n if kind == "Figure" else self.tab_n
        r = self._field(para, f"SEQ {kind} \\* ARABIC", str(n))
        r.bold = True
        self._run(para, f". {text}", bold=False)
        return para

    def figure(self, path: Path, caption: str, interpretation: str | None = None, width_cm: float = 16.4):
        self.fig_n += 1
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        para.add_run().add_picture(str(path), width=Cm(width_cm))
        self.caption("Figure", caption)
        if interpretation:
            self.p(interpretation, bold_lead="Interpretation. ")
        return self.fig_n

    def table(self, df: pd.DataFrame, caption: str, widths_cm: list[float] | None = None, font_size: float = 8.5,
              note: str | None = None, first_col_bold: bool = False):
        self.tab_n += 1
        cap = self.caption("Table", caption)
        cap.paragraph_format.keep_with_next = True
        t = self.doc.add_table(rows=1, cols=len(df.columns))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = "Table Grid"
        self._table_borders(t)
        hdr = t.rows[0]
        self._repeat_header(hdr)
        for i, col in enumerate(df.columns):
            cell = hdr.cells[i]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(col))
            r.bold = True
            r.font.size = Pt(font_size)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            self._shade(cell, HEADER_FILL)
        for ri, row in enumerate(df.itertuples(index=False)):
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""
                para = cells[i].paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                r = para.add_run("" if pd.isna(val) else str(val))
                r.font.size = Pt(font_size)
                if first_col_bold and i == 0:
                    r.bold = True
                if ri % 2 == 1:
                    self._shade(cells[i], BAND_FILL)
        for row in t.rows:
            for c in row.cells:
                c.paragraphs[0].paragraph_format.space_after = Pt(0)
                c.paragraphs[0].paragraph_format.line_spacing = 1.0
        if widths_cm:
            t.autofit = False
            for row in t.rows:
                for i, w in enumerate(widths_cm):
                    row.cells[i].width = Cm(w)
        if note:
            self.p(note, size=8.5, italic=True, color=MUTED, space_after=10)
        else:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    @staticmethod
    def _shade(cell, fill: str):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tcPr.append(shd)

    @staticmethod
    def _repeat_header(row):
        trPr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        trPr.append(el)

    @staticmethod
    def _table_borders(t):
        tblPr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "C9CED6")
            borders.append(el)
        tblPr.append(borders)

    def toc(self, title: str, instr: str, placeholder: str):
        self.p(title, size=16, color=ACCENT).runs[0].bold = True
        para = self.doc.add_paragraph()
        self._field(para, instr, placeholder)

    def landscape(self, on: bool):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        if on:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Mm(297), Mm(210)
        else:
            sec.orientation = WD_ORIENT.PORTRAIT
            sec.page_width, sec.page_height = Mm(210), Mm(297)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, MARGIN)
        sec.different_first_page_header_footer = False
        return sec

    def save(self, path: Path):
        self.doc.core_properties.title = "End-to-End Healthcare Data Analytics"
        self.doc.core_properties.subject = "Diabetic inpatient 30-day readmission analysis"
        self.doc.save(str(path))
